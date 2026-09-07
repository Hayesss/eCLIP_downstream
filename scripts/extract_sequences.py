"""Extract strand-oriented genomic windows with hg38 / GENCODE v38 provenance."""
from pathlib import Path
import argparse
import collections
import csv
import gzip
import hashlib
import json
import mmap
import re
import zipfile
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
COMPLEMENT = str.maketrans('ACGTRYKMSWBDHVN', 'TGCAYRMKSWVHDBN')


def sha256(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        while chunk := stream.read(8 * 1024 * 1024):
            h.update(chunk)
    return h.hexdigest()


def reverse_complement(seq):
    return seq.translate(COMPLEMENT)[::-1]


class Genome:
    def __init__(self, path):
        self.stream = path.open('rb')
        self.data = mmap.mmap(self.stream.fileno(), 0, access=mmap.ACCESS_READ)
        self.index = {}
        fai = path.with_name(path.name + '.fai')
        if fai.exists():
            for line in fai.read_text().splitlines():
                name, length, offset, bases, width = line.split('\t')
                self.index[name] = tuple(map(int, (length, offset, bases, width)))
        else:
            pos = 0
            while pos < len(self.data):
                assert self.data[pos:pos+1] == b'>'
                header_end = self.data.find(b'\n', pos)
                name = self.data[pos+1:header_end].split()[0].decode()
                offset = header_end + 1
                first_end = self.data.find(b'\n', offset)
                width = first_end + 1 - offset
                bases = len(self.data[offset:first_end].rstrip(b'\r'))
                next_header = self.data.find(b'\n>', offset)
                stop = next_header + 1 if next_header >= 0 else len(self.data)
                nlines, rest = divmod(stop-offset, width)
                length = nlines*bases
                if rest:
                    length += len(self.data[stop-rest:stop].rstrip(b'\r\n'))
                self.index[name] = (length, offset, bases, width)
                pos = stop
            fai.write_text(''.join(f'{name}\t' + '\t'.join(map(str, item)) + '\n' for name, item in self.index.items()), encoding='ascii')
        print('REFERENCE_INDEX', len(self.index), 'contigs', flush=True)

    def fetch(self, chrom, start, end):
        length, offset, bases, width = self.index[chrom]
        assert 0 <= start <= end <= length, (chrom, start, end, length)
        left = offset + start//bases*width + start%bases
        right = offset + end//bases*width + end%bases
        seq = self.data[left:right].translate(None, b'\r\n').decode('ascii').upper()
        assert len(seq) == end-start
        return seq


def interval(record, target, genome):
    chrom, start, end, strand = record['chr'], record['start'], record['end'], record['strand']
    chrom_len = genome.index[chrom][0]
    n = end-start
    if target is None:
        return start, end, [(start, end)], 0, 'none'
    assert target >= n and chrom_len >= target
    flank5 = (target-n)//2
    flank3 = target-n-flank5
    left = start - (flank5 if strand == '+' else flank3)
    right = end + (flank3 if strand == '+' else flank5)
    boundary = 'none'
    if chrom == 'chrM' and (left < 0 or right > chrom_len):
        boundary = 'circular_wrap_chrM'
        if left < 0:
            segments = [(chrom_len+left, chrom_len), (0, right)]
        else:
            segments = [(left, chrom_len), (0, right-chrom_len)]
    else:
        if left < 0:
            left, right = 0, target
            boundary = 'shifted_inside_chromosome'
        elif right > chrom_len:
            left, right = chrom_len-target, chrom_len
            boundary = 'shifted_inside_chromosome'
        segments = [(left, right)]
    peak_offset = start-left if strand == '+' else right-end
    return segments[0][0], segments[-1][1], segments, peak_offset, boundary


def read_fasta(path):
    header, parts = None, []
    with path.open(encoding='ascii') as stream:
        for line in stream:
            if line.startswith('>'):
                if header is not None:
                    yield header, ''.join(parts)
                header, parts = line[1:].strip(), []
            else:
                parts.append(line.strip())
    if header is not None:
        yield header, ''.join(parts)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True, help='Source eCLIP XLSX; column headers on row 2')
    parser.add_argument('--sheet', default='Sheet1')
    parser.add_argument('--reference-dir', type=Path, default=ROOT / 'references' / 'gencode_v38')
    parser.add_argument('--metadata-dir', type=Path, help='Defaults to --reference-dir')
    parser.add_argument('--output-dir', type=Path, default=ROOT / 'results' / 'sequences_hg38_gencodev38')
    parser.add_argument('--expected-count', type=int, help='Optional required number of input sites')
    args = parser.parse_args()
    SOURCE = args.input.resolve()
    REFDIR = args.reference_dir.resolve()
    FASTA = REFDIR / 'GRCh38.primary_assembly.genome.fa'
    OUT = args.output_dir.resolve()
    reference = json.loads((REFDIR / 'reference_manifest.json').read_text(encoding='utf-8'))
    assert reference['gencode_release'] == 38
    genome = Genome(FASTA)
    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook[args.sheet]
    headers = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
    records = []
    wanted = ['chr', 'start', 'end', 'name', 'strand', 'gc', 'gene_name', 'gene_id', 'transcript_ids', 'feature_type_top', 'feature_types', 'q_min', 'q_max', 'enrichment_l2or_mean', 'input_sum', 'clip_sum']
    positions = {key: headers.index(key) for key in wanted}
    gc = collections.Counter()
    gc_bad = []
    genes, transcripts = set(), set()
    for row_number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
        record = {key: row[index] for key, index in positions.items()}
        record['source_excel_row'] = row_number
        record['site_id'] = f'SARS1_eCLIP_{row_number-2:06d}'
        chrom, start, end, strand = (record[k] for k in ['chr', 'start', 'end', 'strand'])
        assert isinstance(start, int) and isinstance(end, int) and strand in ['+', '-']
        assert 0 <= start < end <= genome.index[chrom][0]
        sequence = genome.fetch(chrom, start, end)
        assert set(sequence) <= set('ACGTRYKMSWBDHVN')
        observed_gc = (sequence.count('G')+sequence.count('C'))/len(sequence)
        gc['tested'] += 1
        if abs(observed_gc-float(record['gc'])) <= 1e-6:
            gc['bed_0based_halfopen_matches'] += 1
        else:
            gc_bad.append({'site_id': record['site_id'], 'source_gc': record['gc'], 'extracted_gc': observed_gc})
        if start > 0:
            alternative = genome.fetch(chrom, start-1, end)
            alt_gc = (alternative.count('G')+alternative.count('C'))/len(alternative)
            if abs(alt_gc-float(record['gc'])) <= 1e-6:
                gc['alternative_1based_inclusive_matches'] += 1
        record['original_sequence'] = sequence if strand == '+' else reverse_complement(sequence)
        genes.update(re.findall(r'ENSG\d+\.\d+', str(record['gene_id'])))
        transcripts.update(re.findall(r'ENST\d+\.\d+', str(record['transcript_ids'])))
        records.append(record)
    workbook.close()
    print('GC_COORDINATE_CHECK', dict(gc), 'mismatches', gc_bad[:5], flush=True)
    if not records:
        raise ValueError('No binding sites found')
    if args.expected_count is not None and len(records) != args.expected_count:
        raise ValueError(f'Expected {args.expected_count} sites, received {len(records)}')
    assert gc['bed_0based_halfopen_matches']/len(records) >= 0.99, 'Investigate reference or coordinate mismatch before exporting'
    metadata = args.metadata_dir.resolve() if args.metadata_dir else REFDIR
    def ref_ids(kind):
        with gzip.open(metadata / f'gencode.v38.metadata.{kind}_source.gz', 'rt') as stream:
            return {line.split('\t')[0] for line in stream}
    assert genes <= ref_ids('Gene')
    assert transcripts <= ref_ids('Transcript')
    OUT.mkdir(parents=True, exist_ok=True)
    fields = ['site_id','source_excel_row','source_peak_name','gene_name','gene_id','transcript_ids','feature_type_top','feature_types','genome','annotation_release','chr','strand','peak_start0','peak_end0','peak_length_bp','sequence_start0','sequence_end0','genomic_segments_0based_halfopen','sequence_length_bp','peak_offset0_in_sequence','peak_end_offset0_in_sequence','upstream_5prime_bp','downstream_3prime_bp','boundary_handling','N_count','gc_fraction','source_gc','q_min','q_max','enrichment_l2or_mean','input_sum','clip_sum','sequence_5prime_to_3prime_DNA']
    outputs = []
    summary = {'source_file': str(SOURCE), 'source_sha256': sha256(SOURCE), 'source_sheet': args.sheet, 'site_count': len(records), 'assembly': 'hg38 / GRCh38', 'annotation_release': 'GENCODE v38', 'reference': reference, 'coordinate_convention': '0-based, end-exclusive (BED)', 'sequence_orientation': 'strand-specific 5prime-to-3prime DNA alphabet; negative strand reverse complemented', 'window_definition': 'Final total length 500 or 1000 bp, containing the complete peak; flanks balanced in transcript direction with any odd extra base placed on the 3prime flank.', 'sequence_context': 'Contiguous genomic DNA, not a spliced mature transcript; no sample-specific variants applied.', 'gencode_v38_id_validation': {'gene_ids': len(genes), 'transcript_ids': len(transcripts), 'all_matched': True}, 'source_gc_comparison': dict(gc), 'source_gc_mismatches': gc_bad, 'outputs': {}}
    for label, target in [('binding_sites', None), ('centered_500bp', 500), ('centered_1000bp', 1000)]:
        stem = f'SARS1_eCLIP_{label}_hg38_GENCODEv38'
        tsv, fasta = OUT / (stem+'.tsv'), OUT / (stem+'.fasta')
        stats = collections.Counter()
        min_len, max_len = 10**9, 0
        with tsv.open('w', encoding='utf-8-sig', newline='') as table, fasta.open('w', encoding='ascii', newline='\n') as seqfile:
            writer = csv.DictWriter(table, fieldnames=fields, delimiter='\t')
            writer.writeheader()
            for record in records:
                s, e, segments, offset, boundary = interval(record, target, genome)
                seq = ''.join(genome.fetch(record['chr'], a, b) for a, b in segments)
                if record['strand'] == '-':
                    seq = reverse_complement(seq)
                peak_len = record['end']-record['start']
                assert len(seq) == (target or peak_len)
                assert seq[offset:offset+peak_len] == record['original_sequence']
                stats['records'] += 1
                stats['total_bases'] += len(seq)
                stats['records_with_N'] += int('N' in seq)
                stats['N_bases'] += seq.count('N')
                stats[boundary] += 1
                min_len, max_len = min(min_len,len(seq)), max(max_len,len(seq))
                segtext = ';'.join(f"{record['chr']}:{a}-{b}" for a,b in segments)
                result = {key: record[key] for key in ['site_id','source_excel_row','gene_name','gene_id','transcript_ids','feature_type_top','feature_types','chr','strand','q_min','q_max','enrichment_l2or_mean','input_sum','clip_sum']}
                result.update(source_peak_name=record['name'], genome='hg38', annotation_release='GENCODE v38', peak_start0=record['start'], peak_end0=record['end'], peak_length_bp=peak_len, sequence_start0=s, sequence_end0=e, genomic_segments_0based_halfopen=segtext, sequence_length_bp=len(seq), peak_offset0_in_sequence=offset, peak_end_offset0_in_sequence=offset+peak_len, upstream_5prime_bp=offset, downstream_3prime_bp=len(seq)-offset-peak_len, boundary_handling=boundary, N_count=seq.count('N'), gc_fraction=round((seq.count('G')+seq.count('C'))/len(seq),8), source_gc=record['gc'], sequence_5prime_to_3prime_DNA=seq)
                writer.writerow(result)
                safe_gene = re.sub(r'[^A-Za-z0-9_.:+-]', '_', str(record['gene_name']))
                seqfile.write(f">{record['site_id']}|gene={safe_gene}|hg38={segtext}|strand={record['strand']}|length={len(seq)}|peak_offset0={offset}|source_row={record['source_excel_row']}\n")
                for start in range(0,len(seq),80):
                    seqfile.write(seq[start:start+80]+'\n')
        # Verify delivered FASTA and table independently agree for every record.
        count = 0
        fasta_iter = iter(read_fasta(fasta))
        with tsv.open(encoding='utf-8-sig', newline='') as table:
            for result in csv.DictReader(table, delimiter='\t'):
                header, seq = next(fasta_iter)
                assert header.split('|')[0] == result['site_id']
                assert seq == result['sequence_5prime_to_3prime_DNA']
                assert len(seq) == int(result['sequence_length_bp'])
                count += 1
        assert count == len(records) and next(fasta_iter,None) is None
        stats.update({'min_length_bp':min_len,'max_length_bp':max_len,'fasta_tsv_verified_records':count})
        summary['outputs'][label] = dict(stats)
        outputs.extend([tsv,fasta])
        print('EXPORTED', label, dict(stats), flush=True)
    summary['file_sha256'] = {p.name:sha256(p) for p in outputs}
    summary_path = OUT / 'sequence_extraction_QC.json'
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    readme = OUT / 'README_序列与坐标说明.md'
    readme.write_text('''# SARS1 eCLIP 结合区间及扩展序列

原始结合区间、以结合区间为中心的总长 500 bp 和总长 1000 bp，各提供一份 FASTA 和一份 TSV。每份包含全部 50,613 个位点，沿用原表顺序，不做合并或进一步显著性筛选。

## 固定版本与序列含义

- 坐标体系：hg38 / GRCh38。
- 注释版本：GENCODE v38；保留原表结合位点的基因、转录本和区域注释，所有带版本基因/转录本 ID 已与 v38 官方清单核对。该步骤不是重新计算外显子或 UTR 注释。
- 基因组 FASTA 来自 GENCODE release_38 的 GRCh38.primary_assembly.genome.fa.gz，官方 MD5 已核对；来源、校验值及输出校验值记录在 QC JSON 中。
- 所有序列为参考基因组序列，以 DNA 字母 A/C/G/T 表示。正链取原序列，负链取反向互补，均为对应 RNA 方向的 5′→3′。未加入样本特异变异。
- 提取的是完整 eCLIP 结合区间，不是对最小结合基序或单碱基接触点的实验确认。
- 扩展按连续基因组区间进行，可能包含内含子或基因间区域；没有按成熟转录本拼接。如果用于成熟 mRNA 构建设计，需要进一步选定转录本并处理剪接。

## 长度与坐标

- 500/1000 bp 是扩展后包括结合区间在内的总长度，并非每侧各加 500/1000 bp。
- 扩展完整保留原始结合区间，5′和3′两侧尽量均分；无法均分时，多出的 1 bp 分配到3′侧。
- 表中所有带 start0/end0 或 offset0 的坐标均为 0-based、左闭右开；将普通单段区间用于浏览器的 1-based 显示时，start0 加 1，end0 不变。
- peak_offset0_in_sequence / peak_end_offset0_in_sequence 标明原始结合区间在输出序列中的位置；Python 切片 sequence[start:end] 可精确取回原始结合序列。
- 线性染色体边缘若越界，向染色体内部平移以保持总长度，并用 boundary_handling 标记。chrM 按环状序列跨原点拼接；此时 sequence_start0 可以大于 sequence_end0，以 genomic_segments_0based_halfopen 为准确区间说明。
- genomic_segments 列按基因组正向拼接顺序列出；strand 为负时，最终序列再对整个拼接结果反向互补。
- N_count 标出参考序列未知碱基数，没有自动丢弃或替换 N。

## 使用文件

- FASTA：用于序列分析工具。每条序列标题含稳定 site_id、基因名、坐标、链及原表行号。
- TSV：可通过 Excel 的“数据→从文本/CSV”导入，分隔符选择制表符。建议将基因名和 ID 列按文本导入，避免 Excel 自动转换。
- source_excel_row：原始 Sheet1 中的实际行号，可直接追溯。
- annotation_release 只指本次结合位点注释约定；这些文件不使用原表 uga_* 列确定结合序列。
- QC JSON：记录区间长度、GC 比对、边界处理、FASTA/TSV 全量一致性校验及文件 SHA-256。

公开参考：https://www.gencodegenes.org/human/release_38.html
'''.replace('50,613', f'{len(records):,}'), encoding='utf-8')
    outputs.extend([readme,summary_path])
    bundle = OUT / 'SARS1_eCLIP_sequences_hg38_GENCODEv38_all.zip'
    with zipfile.ZipFile(bundle, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for path in outputs:
            archive.write(path, path.name)
    print('COMPLETE', str(bundle), bundle.stat().st_size, flush=True)


if __name__ == '__main__':
    if not __debug__:
        raise SystemExit('Run without -O: scientific validation assertions must remain enabled.')
    main()
