"""Compare versioned eCLIP gene/transcript IDs against public GENCODE releases."""
import argparse
import csv
import gzip
import hashlib
import json
from pathlib import Path
import re
import urllib.request
import openpyxl


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--sheet', default='Sheet1')
    parser.add_argument('--releases', nargs='+', type=int, default=[29,37,38,39,40,41,42,43])
    parser.add_argument('--cache-dir', type=Path, default=Path('references')/'gencode_metadata')
    parser.add_argument('--output-dir', type=Path, default=Path('results')/'reference_audit')
    args = parser.parse_args()
    if any(v < 1 for v in args.releases):
        parser.error('Release numbers must be positive')
    args.cache_dir.mkdir(parents=True, exist_ok=True)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    sheet = workbook[args.sheet]
    headers = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
    identifiers = {key:set() for key in ['gene_id','transcript_ids','uga_gene_id']}
    positions = {key:headers.index(key) for key in identifiers}
    for row in sheet.iter_rows(min_row=3, values_only=True):
        for key, index in positions.items():
            prefix = 'ENST' if key == 'transcript_ids' else 'ENSG'
            identifiers[key].update(re.findall(prefix+r'\d+\.\d+', str(row[index])))
    workbook.close()
    rows, sources = [], []
    for release in args.releases:
        for kind in ['Gene','Transcript']:
            name = f'gencode.v{release}.metadata.{kind}_source.gz'
            url = f'https://ftp.ebi.ac.uk/pub/databases/gencode/Gencode_human/release_{release}/{name}'
            path = args.cache_dir/name
            if not path.exists():
                with urllib.request.urlopen(url, timeout=60) as source:
                    payload = source.read()
                gzip.decompress(payload)  # Verify before accepting into the cache.
                temporary = path.with_name(path.name+'.partial')
                temporary.write_bytes(payload)
                temporary.replace(path)
            payload = path.read_bytes()
            reference_ids = {line.split('\t')[0] for line in gzip.decompress(payload).decode().splitlines()}
            sources.append({'url':url,'sha256':hashlib.sha256(payload).hexdigest()})
            for key, values in identifiers.items():
                if (key == 'transcript_ids') != (kind == 'Transcript'):
                    continue
                rows.append({'release':release,'column':key,'unique_ids':len(values),'matched_ids':len(values & reference_ids),'missing_ids':len(values-reference_ids)})
                print(rows[-1], flush=True)
    with (args.output_dir/'gencode_version_comparison.tsv').open('w', encoding='utf-8', newline='') as out:
        writer = csv.DictWriter(out, fieldnames=list(rows[0]), delimiter='\t')
        writer.writeheader()
        writer.writerows(rows)
    report = {'source_file':str(args.input.resolve()),'source_sha256':hashlib.sha256(args.input.read_bytes()).hexdigest(),'results':rows,'public_sources':sources,'interpretation':'Exact ID membership is evidence of release compatibility, not direct confirmation of the original workflow or feature coordinates.'}
    (args.output_dir/'gencode_version_comparison.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')


if __name__ == '__main__':
    main()
